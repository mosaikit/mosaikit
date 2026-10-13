#!/usr/bin/env python3
# SPDX-FileCopyrightText: 2026 Massimo Antonini
# SPDX-License-Identifier: MPL-2.0
"""
Builds the documents of a release from the Markdown and YAML of the repository (ADR-0023):
Word and PDF guides, an Excel workbook of compliance and requirements, a PowerPoint overview of
the release, extended release notes, and a zip of all of them with their SHA-256 checksums.

    python3 docs/build/build.py --version 0.2.0 --output target/docs [--no-pdf]

Needs pandoc (2.9 or later), xelatex for the PDFs, and the Python packages openpyxl and PyYAML.
The Markdown files stay the source: nothing here is edited by hand.
"""
from __future__ import annotations

import argparse
import datetime
import hashlib
import os
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
FILTER = Path(__file__).with_name("links.lua")
PDF_HEADER = Path(__file__).with_name("pdf-header.tex")
SKIPPED = {"node_modules", "target", "dist", "coverage"}

GUIDES = {
    "User-Guide": ("User and administrator guide", "guide",
                   ["README", "installation", "using", "administration", "extensions", "configuration", "operations"]),
    "Developer-Guide": ("Developer guide", "development",
                        ["README", "architecture", "development", "extension-development", "api", "versions", "release", "roadmap"]),
    "Compliance": ("Compliance: AgID, ACN QC2 and GDPR", "compliance",
                   ["README", "matrix", "non-conformities", "evidence", "secure-development"]),
}

LINK = re.compile(r"\[([^\]]*)\]\([^)]*\)")
CODE = re.compile(r"`([^`]*)`")
EMPHASIS = re.compile(r"\*\*?([^*]+)\*\*?")


def plain(cell: str) -> str:
    """The text of a Markdown table cell, without links, code marks and emphasis."""
    return EMPHASIS.sub(r"\1", CODE.sub(r"\1", LINK.sub(r"\1", cell))).replace("<br>", "\n").strip()


def tables(markdown: str) -> list[tuple[str, list[str], list[list[str]]]]:
    """The pipe tables of a Markdown text, each with the heading above it, its header and rows."""
    found = []
    heading = ""
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("#"):
            heading = line.lstrip("#").strip()
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1]):
            header = [plain(cell) for cell in line.strip("|").split("|")]
            rows = []
            i += 2
            while i < len(lines) and lines[i].startswith("|"):
                cells = [plain(cell) for cell in lines[i].strip("|").split("|")]
                rows.append((cells + [""] * len(header))[: len(header)])
                i += 1
            found.append((heading, header, rows))
            continue
        i += 1
    return found


def pandoc(sources: list[Path], output: Path, metadata: dict[str, str], extra: list[str] = ()) -> None:
    # Pandoc Markdown reads the GitHub Markdown of the docs, and gives pipe tables relative widths
    # so that long cells wrap in PDF and Word.
    command = ["pandoc", "--from", "markdown-yaml_metadata_block+autolink_bare_uris", "--columns", "80",
               "--lua-filter", str(FILTER), "--toc", "--standalone", "-o", str(output), *extra]
    for key, value in metadata.items():
        command += ["--metadata", f"{key}={value}"]
    subprocess.run(command + [str(source) for source in sources], check=True, cwd=ROOT)


def guides(version: str, output: Path, pdf: bool) -> list[Path]:
    built = []
    date = datetime.date.today().isoformat()
    for name, (title, folder, pages) in GUIDES.items():
        sources = [DOCS / folder / f"{page}.md" for page in pages]
        missing = [str(source) for source in sources if not source.exists()]
        if missing:
            raise SystemExit(f"Missing pages of the {title}: {missing}")
        metadata = {"title": f"Mosaikit — {title}", "subtitle": f"Version {version}", "date": date}
        docx = output / f"Mosaikit-{name}-{version}.docx"
        pandoc(sources, docx, metadata)
        built.append(docx)
        if pdf:
            target = output / f"Mosaikit-{name}-{version}.pdf"
            pandoc(sources, target, metadata, ["--pdf-engine=xelatex", "-V", "geometry:margin=2cm",
                                               "-V", "fontsize=10pt", "-V", "colorlinks=true",
                                               "-V", "mainfont=DejaVu Sans", "-V", "monofont=DejaVu Sans Mono",
                                               "--include-in-header", str(PDF_HEADER)])
            built.append(target)
    adrs = sorted(path for path in (DOCS / "adr").glob("[0-9]*.md"))
    decisions = output / f"Mosaikit-Architecture-Decisions-{version}.docx"
    pandoc(adrs, decisions, {"title": "Mosaikit — Architecture decision records",
                             "subtitle": f"Version {version}", "date": date})
    built.append(decisions)
    return built


def requirement_tests() -> dict[str, list[str]]:
    """The test files that reference each requirement, as the requirements test checks them."""
    references: dict[str, list[str]] = {}
    for directory, children, files in os.walk(ROOT):
        children[:] = [child for child in children if child not in SKIPPED and not child.startswith(".")]
        if "test" not in Path(directory).relative_to(ROOT).parts:
            continue
        for name in files:
            if name.endswith(("Test.java", "IT.java", ".test.ts")):
                text = (Path(directory) / name).read_text(encoding="utf-8")
                for requirement in set(re.findall(r"MK-\d{3}", text)):
                    references.setdefault(requirement, []).append(name)
    return references


def workbook(version: str, output: Path) -> Path:
    book = Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")

    def sheet(title: str, header: list[str], rows: list[list[str]]) -> None:
        ws = book.create_sheet(title)
        ws.append(header)
        for row in rows:
            ws.append(row)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        for column, name in enumerate(header, start=1):
            width = max([len(str(name))] + [min(len(str(row[column - 1])), 80) for row in rows]) + 2
            ws.column_dimensions[get_column_letter(column)].width = min(width, 80)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    matrix_rows = []
    matrix_header = None
    for heading, header, rows in tables((DOCS / "compliance" / "matrix.md").read_text(encoding="utf-8")):
        if header and header[0] == "ID":
            matrix_header = ["Area"] + header
            matrix_rows += [[heading] + row for row in rows]
    if matrix_header:
        sheet("Compliance matrix", matrix_header, matrix_rows)
    nc = tables((DOCS / "compliance" / "non-conformities.md").read_text(encoding="utf-8"))
    if nc:
        sheet("Non-conformities", nc[0][1], nc[0][2])
    tests = requirement_tests()
    requirements = []
    for file in sorted((DOCS / "requirements").glob("MK-*.yaml")):
        data = yaml.safe_load(file.read_text(encoding="utf-8"))
        requirements.append([data["id"], data["title"], data["area"], data["priority"], data["status"],
                             " ".join(data.get("description", "").split()),
                             "\n".join(data.get("acceptance", [])),
                             ", ".join(data.get("depends", []) or []),
                             ", ".join(sorted(set(tests.get(data["id"], [])))) ])
    sheet("Requirements", ["ID", "Title", "Area", "Priority", "Status", "Description", "Acceptance",
                           "Depends on", "Verified by"], requirements)
    del book["Sheet"]
    target = output / f"Mosaikit-Compliance-{version}.xlsx"
    book.properties.title = f"Mosaikit {version}: compliance and requirements"
    book.save(target)
    return target


def changelog_section(version: str) -> str:
    text = (ROOT / "CHANGELOG.md").read_text(encoding="utf-8")
    for heading in (f"## [{version}]", "## [Unreleased]"):
        start = text.find(heading)
        if start >= 0:
            body = text[text.index("\n", start) + 1:]
            end = body.find("\n## [")
            return (body if end < 0 else body[:end]).strip()
    return ""


def overview(version: str) -> str:
    """The Markdown of the release overview: what is new, requirements and compliance."""
    changes = changelog_section(version)
    requirements = [yaml.safe_load(file.read_text(encoding="utf-8"))
                    for file in sorted((DOCS / "requirements").glob("MK-*.yaml"))]
    done = [r for r in requirements if r["status"] == "done"]
    statuses: dict[str, int] = {}
    for _heading, header, rows in tables((DOCS / "compliance" / "matrix.md").read_text(encoding="utf-8")):
        if header and header[0] == "ID" and "Status" in header:
            for row in rows:
                statuses[row[header.index("Status")]] = statuses.get(row[header.index("Status")], 0) + 1
    open_nc = []
    nc = tables((DOCS / "compliance" / "non-conformities.md").read_text(encoding="utf-8"))
    if nc:
        header, rows = nc[0][1], nc[0][2]
        state = header.index("Status") if "Status" in header else len(header) - 1
        open_nc = [row for row in rows if not row[state].startswith("Closed")]

    slides = [f"% Mosaikit {version}\n% Release overview\n% {datetime.date.today().isoformat()}\n"]
    for block in re.split(r"^### ", changes, flags=re.M):
        block = block.strip()
        if not block or block.startswith("-"):
            continue
        kind, _, rest = block.partition("\n")
        items = [re.sub(r"\s+", " ", item).strip() for item in re.split(r"^- ", rest, flags=re.M) if item.strip()]
        for chunk in range(0, len(items), 5):
            bullets = "\n".join(f"- {plain(item.split(': ')[0].split(' (')[0])}" for item in items[chunk:chunk + 5])
            slides.append(f"# {kind}\n\n{bullets}\n")
    slides.append(f"# Requirements\n\n{len(done)} of {len(requirements)} requirements are done.\n\n"
                  + "\n".join(f"- {r['id']} {r['title']}" for r in done[-8:]) + "\n")
    slides.append("# Compliance\n\n| Status | Measures |\n|---|---|\n"
                  + "\n".join(f"| {status} | {count} |" for status, count in sorted(statuses.items())) + "\n")
    if open_nc:
        for chunk in range(0, len(open_nc), 6):
            slides.append("# Open non-conformities\n\n"
                          + "\n".join(f"- {row[0]} {row[1]} ({row[2]})" for row in open_nc[chunk:chunk + 6]) + "\n")
    slides.append("# Get it\n\n- Portable archives for Windows, Linux and macOS\n- Container image, Docker Compose"
                  " and Helm chart\n- Extensions as a signed catalog for the marketplace\n- Guides, compliance"
                  " workbook and this overview in the documents bundle\n")
    return "\n\n".join(slides)


def release_notes(version: str) -> str:
    """The CHANGELOG section, with the state of requirements and non-conformities."""
    requirements = [yaml.safe_load(file.read_text(encoding="utf-8"))
                    for file in sorted((DOCS / "requirements").glob("MK-*.yaml"))]
    nc = tables((DOCS / "compliance" / "non-conformities.md").read_text(encoding="utf-8"))
    lines = [f"# Mosaikit {version}", "", changelog_section(version), "", "## Requirements", "",
             "| ID | Title | Status |", "|---|---|---|"]
    lines += [f"| {r['id']} | {r['title']} | {r['status']} |" for r in requirements]
    if nc:
        header, rows = nc[0][1], nc[0][2]
        lines += ["", "## Open non-conformities", "", "| " + " | ".join(header) + " |",
                  "|" + "---|" * len(header)]
        lines += ["| " + " | ".join(row) + " |" for row in rows if not row[-1].startswith("Closed")]
    return "\n".join(lines) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--version", required=True)
    parser.add_argument("--output", type=Path, default=ROOT / "target" / "docs")
    parser.add_argument("--no-pdf", action="store_true", help="skip the PDFs, which need xelatex")
    args = parser.parse_args()
    output: Path = args.output.resolve()
    if output.exists():
        shutil.rmtree(output)
    output.mkdir(parents=True)

    built = guides(args.version, output, not args.no_pdf)
    built.append(workbook(args.version, output))
    notes = output / f"Mosaikit-Release-Notes-{args.version}.md"
    notes.write_text(release_notes(args.version), encoding="utf-8")
    built.append(notes)
    slides = output / "overview.md"
    slides.write_text(overview(args.version), encoding="utf-8")
    deck = output / f"Mosaikit-Release-Overview-{args.version}.pptx"
    subprocess.run(["pandoc", "--from", "markdown", "-o", str(deck), str(slides)], check=True)
    slides.unlink()
    built.append(deck)

    sums = output / "SHA256SUMS"
    sums.write_text("".join(f"{hashlib.sha256(path.read_bytes()).hexdigest()}  {path.name}\n" for path in built))
    bundle = output / f"mosaikit-docs-{args.version}.zip"
    with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as archive:
        for path in built + [sums]:
            archive.write(path, path.name)
    for path in built + [sums, bundle]:
        print(f"{path.stat().st_size:>10}  {path.relative_to(ROOT) if path.is_relative_to(ROOT) else path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
